"""Маленькие учебные файлы для тестов Ingestion. Создаются на лету — бинарники в git не нужны."""

from __future__ import annotations

import io
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


def png_bytes() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (40, 30), "steelblue").save(buf, format="PNG")
    return buf.getvalue()


def make_pptx(path: Path) -> Path:
    """4 слайда: теория с кодом и таблицей, продолжение темы, пример, слайд-картинка."""
    prs = Presentation()

    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Loops in Python"
    body = s1.placeholders[1].text_frame
    body.text = "A loop repeats a block of code. LO 11.2.1.1"
    sub = body.add_paragraph()
    sub.text = "for — a fixed number of times"
    sub.level = 1
    code = s1.shapes.add_textbox(Inches(1), Inches(4), Inches(5), Inches(1)).text_frame
    code.text = "for i in range(3):"
    code.paragraphs[0].runs[0].font.name = "Consolas"
    run = code.add_paragraph().add_run()
    run.text = "    print(i)"
    run.font.name = "Consolas"
    table = s1.shapes.add_table(2, 2, Inches(1), Inches(5.5), Inches(4), Inches(1)).table
    for (r, c), text in {(0, 0): "Op", (0, 1): "Meaning", (1, 0): "%", (1, 1): "remainder"}.items():
        table.cell(r, c).text = text
    s1.notes_slide.notes_text_frame.text = "Explain range() carefully."

    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Loops in Python"
    s2.placeholders[1].text_frame.text = "while — until a condition becomes false."

    s3 = prs.slides.add_slide(prs.slide_layouts[1])
    s3.shapes.title.text = "Practice"
    s3.placeholders[1].text_frame.text = "Example: sum of a list"
    plain = s3.shapes.add_textbox(Inches(1), Inches(4), Inches(5), Inches(1)).text_frame
    plain.text = "total = 0\nfor x in numbers:\n    total += x"   # код без моноширинного шрифта

    s4 = prs.slides.add_slide(prs.slide_layouts[6])            # пустой макет
    s4.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(1))

    prs.save(str(path))
    return path


def make_pdf(path: Path, pages: int = 4, scanned: bool = True) -> Path:
    """Учебник с оглавлением, кодом и колонтитулом; последняя страница — «скан»."""
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"Chapter {i + 1}. Data types", fontsize=20)
        page.insert_text(
            (72, 110),
            "A variable is a named storage location in memory. Its value can change.",
            fontsize=11,
        )
        page.insert_text(
            (72, 140), "def square(x):\n    return x * x", fontname="cour", fontsize=10
        )
        page.insert_text((72, 790), "NIS Informatics 11 - textbook", fontsize=8)
        page.insert_text((300, 810), str(i + 1), fontsize=8)
    if scanned:
        page = doc.new_page()
        page.insert_image(pymupdf.Rect(72, 72, 400, 400), stream=png_bytes())
    doc.set_toc([[1, "Unit 1. Basics", 1], [2, "1.1 Variables", 2], [1, "Unit 2. Types", 3]])
    doc.save(str(path))
    return path


def make_objectives_xlsx(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["Code", "Learning objective", "Term"])
    ws.append(["11.2A Programming"])                            # строка раздела
    ws.append(["11.2.1.1", "Use loops to solve problems", 1])
    ws.append(["11.2.1.2", "Write functions with parameters", "Term 2"])
    ws.append(["11.3B Databases"])
    ws.append(["11.3.1.1", "Explain the purpose of a primary key", "III"])
    ws.append(["11.3.1.2", None, 3])                            # без формулировки
    wb.save(path)
    return path


def make_objectives_csv(path: Path) -> Path:
    text = (
        "Код;Цель обучения;Четверть\n"
        "11.1.1.1;Объяснять назначение операционной системы;1\n"
        "11.1.1.2;Сравнивать типы памяти компьютера;1\n"
    )
    path.write_bytes(text.encode("cp1251"))        # как сохраняет Excel на русской Windows
    return path


def make_objectives_docx(path: Path) -> Path:
    doc = DocxDocument()
    doc.add_paragraph("Unit 11.4 Networks")
    doc.add_paragraph("11.4.1.1 Describe the layers of the TCP/IP model and")
    doc.add_paragraph("the role of each layer")
    doc.add_paragraph("")
    doc.add_paragraph("11.4.1.2 Explain how DNS works")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Code"
    table.cell(0, 1).text = "Objective"
    table.cell(1, 0).text = "11.4.2.1"
    table.cell(1, 1).text = "Compare IPv4 and IPv6 addressing"
    doc.save(str(path))
    return path


def make_html_article(path: Path) -> Path:
    """Статья в духе bzfar.org: меню, коды целей, подзаголовки жирным, код, таблица."""
    path.write_text(
        """<html><head><meta charset="utf-8"><title>Loops in Python - Lessons - BZFAR</title></head>
<body>
<nav><a href="/">Home</a> Menu that must be skipped</nav>
<table><tr><td class="eText">
12.5.3.5 describe a syntax error in a program code<br>
<p><b>Types of loops</b></p>
<p>A loop repeats a block of code. Writing <code>whlie</code> instead of
<code>while</code> is a syntax error.</p><!--<s5174>-->
<span>Loops save time.</span><!-- uCoz comment -->
<div><p>Nested block.</p></div>
<p><code>flag = True<br>
while flag:<br>
&nbsp;&nbsp;&nbsp;&nbsp;answer = input()<br>
&nbsp;&nbsp;&nbsp;&nbsp;flag = answer != "q"</code></p>
<ul><li>for — a fixed number of times</li><li>while — until a condition is false</li></ul>
<h3>Example</h3>
<pre>for i in range(3):
    print(i)</pre>
<table><tr><th>Loop</th><th>When</th></tr><tr><td>for</td><td>known count</td></tr></table>
</td></tr></table>
<footer>Copyright footer</footer>
<script>var tracking = 1;</script>
</body></html>""",
        encoding="utf-8",
    )
    return path


def make_objectives_html(path: Path) -> Path:
    """Список целей как на странице «CS. Grade 12» сайта bzfar.org."""
    path.write_text(
        """<html><head><meta charset="utf-8"><title>Computer Science 12</title></head>
<body><div class="eText">
<p>Unit 12.1B - Programming paradigms</p>
<p>12.1B Expert systems</p>
<p>12.5.1.2 create a simple expert system</p>
<p>Theory and Tasks 2025</p><p>"Expert systems"</p>
<p>Unit 12.2A - Information security</p>
<p>12.1.2.1 describe protection measures for data and computer systems<br>
such as firewalls, encryption, biometrics</p>
<p>Theory &amp; Exercises 2020</p>
</div></body></html>""",
        encoding="utf-8",
    )
    return path


def make_objectives_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text(
        (72, 72),
        "Section 11.5 Algorithms\n"
        "11.5.1.1 Trace a binary search\n"
        "on a sorted array\n"
        "11.5.1.2 Compare bubble sort and merge sort",
        fontsize=11,
    )
    doc.save(str(path))
    return path
