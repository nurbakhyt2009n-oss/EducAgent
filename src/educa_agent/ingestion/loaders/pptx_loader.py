"""PPTX: по каждому слайду — заголовок, текст фигур в порядке чтения, таблицы, код, заметки.

Код распознаётся по моноширинному шрифту (так его обычно вставляют учителя),
а если шрифт не задан — по виду строк.
"""

from __future__ import annotations

from collections.abc import Iterable, Iterator
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from educa_agent.ingestion.loaders.base import LoadedDocument, RawBlock
from educa_agent.ingestion.normalize import looks_like_code, normalize_markdown

MONO_FONTS = (
    "consolas", "courier", "lucida console", "monaco", "menlo", "source code",
    "jetbrains mono", "fira code", "fira mono", "cascadia", "roboto mono",
    "ubuntu mono", "dejavu sans mono", "inconsolata", "pt mono",
)


class PptxLoader:
    def supports(self, path: Path) -> bool:
        return path.suffix.lower() == ".pptx"

    def load(self, path: Path) -> LoadedDocument:
        prs = Presentation(str(path))
        blocks: list[RawBlock] = []
        image_only: list[int] = []

        for n, slide in enumerate(prs.slides, start=1):
            title_shape = slide.shapes.title
            title = ""
            if title_shape is not None and title_shape.has_text_frame:
                title = " ".join(title_shape.text_frame.text.split())
            section = (title,) if title else ()

            content: list[RawBlock] = []
            visuals = 0
            for shape in _reading_order(_flatten(slide.shapes)):
                if title_shape is not None and shape.shape_id == title_shape.shape_id:
                    continue
                if getattr(shape, "has_table", False):
                    table = _table_markdown(shape.table)
                    if table:
                        content.append(RawBlock(table, "table", n, section))
                elif shape.has_text_frame:
                    block = _text_block(shape, n, section)
                    if block is not None:
                        content.append(block)
                else:
                    visuals += 1          # картинка, диаграмма, видео

            if slide.has_notes_slide:
                notes = normalize_markdown(_plain(slide.notes_slide.notes_text_frame.text))
                if notes:
                    content.append(RawBlock(notes, "notes", n, section))

            if title:
                blocks.append(RawBlock(title, "heading", n, section))
            blocks.extend(content)
            if not content and visuals:
                image_only.append(n)

        return LoadedDocument(
            path=path, kind="pptx", unit="slide", blocks=blocks, image_only_pages=image_only
        )


def _flatten(shapes: Iterable[Any]) -> Iterator[Any]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _flatten(shape.shapes)
        else:
            yield shape


def _reading_order(shapes: Iterable[Any]) -> list[Any]:
    """Сверху вниз, слева направо. У наследованных плейсхолдеров координат может не быть."""
    return sorted(shapes, key=lambda s: (s.top or 0, s.left or 0))


def _plain(text: str) -> str:
    return text.replace("\x0b", "\n")   # <a:br/> python-pptx отдаёт как вертикальную табуляцию


def _text_block(shape: Any, page: int, section: tuple[str, ...]) -> RawBlock | None:
    paragraphs = shape.text_frame.paragraphs
    raw_lines = [_plain(p.text) for p in paragraphs]
    if not any(line.strip() for line in raw_lines):
        return None

    total = mono = 0
    for p in paragraphs:
        for run in p.runs:
            size = len(run.text.strip())
            total += size
            if run.font.name and run.font.name.lower().startswith(MONO_FONTS):
                mono += size
    code_text = "\n".join(raw_lines).strip("\n")
    if (total and mono / total >= 0.6) or looks_like_code(code_text):
        return RawBlock(code_text, "code", page, section)

    if shape.is_placeholder:            # тело слайда — маркированный список
        lines = [
            "  " * p.level + "- " + _plain(p.text).strip()
            for p in paragraphs
            if p.text.strip()
        ]
    else:
        lines = raw_lines
    text = normalize_markdown("\n".join(lines))
    return RawBlock(text, "text", page, section) if text else None


def _table_markdown(table: Any) -> str:
    rows = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()).replace("|", "\\|") for cell in row.cells]
        if any(cells):
            rows.append("| " + " | ".join(cells) + " |")
    if not rows:
        return ""
    width = rows[0].count(" | ") + 1
    rows.insert(1, "|" + " --- |" * width)
    return "\n".join(rows)
